#!/usr/bin/env python3
"""Read original workbooks. Preserve source coordinates and merged header paths.

No guessed Cartesian products: one record per populated, explicitly headed cell.
The JSON intermediate is deterministic and is archived by the Node adapter.
"""
import base64, gzip, hashlib, io, json, re, sys
from pathlib import Path
import openpyxl

ROOT = Path(sys.argv[1] if len(sys.argv)>1 else '.')
OUT = Path(sys.argv[2] if len(sys.argv)>2 else '/tmp/atlas-time-use-extracted.json')
def clean(v): return re.sub(r'\s+', ' ', str(v if v is not None else '')).strip()
def numeric(v):
    if isinstance(v, (int,float)) and not isinstance(v,bool): return v
    s=clean(v)
    return float(s) if re.fullmatch(r'-?\d+(?:\.\d+)?',s) else None
def context(text, ctx):
    t=clean(text).lower()
    if 'all-india' in t or t=='all india': ctx['state']='All India'
    m=re.search(r'\b(rural\+urban|rural|urban)\b',t)
    if m: ctx['sector']={'rural+urban':'Combined','rural':'Rural','urban':'Urban'}[m[1]]
    genders=re.findall(r'\b(male|female|persons?)\b',t)
    if genders: ctx['gender']='Person' if genders[-1].startswith('person') else genders[-1].capitalize()
    for token,key in [('type of day','day_type'),('broad principal activity status','broad_status'),('age group','age_group')]:
        m=re.search(token+r':\s*(.+)$',t)
        if m: ctx[key]=clean(m[1])
    return ctx

result={'schemaVersion':1,'tables':[]}
for file in sorted((ROOT/'data/raw/time-use').glob('report-2024-*.json.gz')):
    source=json.load(gzip.open(file))
    if 'bodyBase64' not in source: continue
    body=base64.b64decode(source['bodyBase64'])
    if hashlib.sha256(body).hexdigest()!=source['sha256']: raise ValueError('Original hash mismatch '+str(file))
    book=openpyxl.load_workbook(io.BytesIO(body),data_only=True,read_only=False)
    table={'id':source['id'],'sourceUrl':source['url'],'sourceSha256':source['sha256'],'documentSha256':source['documentSha256'],'rawFile':str(file.relative_to(ROOT)),'records':[],'sheets':[]}
    for sheet in book:
        rows=list(sheet.values)
        title=clean(rows[0][0]); table.setdefault('title',title)
        # Expand actual merged ranges only, never carry arbitrary neighbouring cells.
        merged={}
        for area in sheet.merged_cells.ranges:
            val=clean(sheet.cell(area.min_row,area.min_col).value)
            for r in range(area.min_row,area.max_row+1):
                for c in range(area.min_col,area.max_col+1): merged[r,c]=val
        def at(r,c): return merged.get((r,c),clean(sheet.cell(r,c).value))
        ctx={'state':'All India'}; width=0; headers=[]; roles={}; start=2
        for ri,row in enumerate(rows,1):
            first=clean(row[0]); values=[clean(x) for x in row]
            if re.match(r'^Table\s*\(',first,re.I):
                start=ri+1; width=0
                for ci,v in enumerate(row,1):
                    if v is not None:roles[ri,ci]='title'
                continue
            enum=re.fullmatch(r'\(1\)',first)
            if enum:
                # Some official headings omit their closing parenthesis (e.g. t2 (23).
                # The consecutive printed identifiers still attest every column.
                cols=[i for i,v in enumerate(values) if re.fullmatch(r'\(\d+\)?',v)]
                if cols!=list(range(len(cols))) or len(cols)<2:raise ValueError(f'Non-contiguous column identifiers {table["id"]}:{ri}')
                if [int(re.search(r'\d+',values[i])[0]) for i in cols]!=list(range(1,len(cols)+1)):raise ValueError('Source column order is not consecutive')
                width=len(cols); heads=[]
                for hr in range(start,ri):
                    rr=rows[hr-1]
                    if any(clean(x) for x in rr[1:width]):heads.append(hr)
                headers=[]
                for ci in range(2,width+1):
                    path=[]
                    for hr in heads:
                        v=at(hr,ci)
                        if v and (not path or path[-1]!=v):path.append(v)
                    if not path:raise ValueError(f'No source header {table["id"]}:{ri}:{ci}')
                    headers.append(path)
                for hr in range(start,ri+1):
                    for ci,v in enumerate(rows[hr-1],1):
                        if v is not None:roles[hr,ci]='header'
                continue
            # Context lines have no values in the numbered data columns.
            data=row[1:width] if width else []
            if first and not any(clean(v) not in ['', '-'] for v in (row[1:width] if width else row[1:])):
                if re.search(r'\b(rural|urban|male|female|persons?)\b',first,re.I) or 'all-india' in first.lower():
                    ctx=context(first,ctx)
                elif width and not first.lower().startswith(('note','*','@')):
                    # State panels in the RSE workbooks have a single geography heading.
                    if table['id'].endswith('_state'):ctx['state']=first
                roles[ri,1]='context'
                if not width:start=ri+1
                continue
            if not width or not first or first in ['-','–']:continue
            if not any(numeric(v) is not None or clean(v)=='-' for v in data):continue
            if re.match(r'^\((?:i|ii|iii)\)|^(?:Note|Source|\*)',first,re.I):continue
            for ci,v in enumerate(data,2):
                raw=clean(v); value=numeric(v)
                if value is None and raw!='-':
                    if not raw:continue
                    raise ValueError(f'Unrecognized data marker {table["id"]}:{ri}:{ci}: {raw}')
                path=headers[ci-2]; text=' / '.join(path)
                is_rse=table['id'].startswith('rse')
                if is_rse: unit='pct_rse';measure='Relative standard error'
                elif table['id']=='t1':unit='count';measure=text
                elif table['id'] in ['t2','t3','t45','t46'] or title.startswith('Table (39)'):unit='pct';measure='Published percentage'
                elif 'percentage of persons' in text.lower():unit='pct';measure='Participation rate'
                elif 'per participant' in text.lower():unit='minutes';measure='Minutes per participant'
                elif 'per person' in text.lower():unit='minutes';measure='Minutes per person'
                elif 'Minutes' in title:unit='minutes';measure='Minutes per person' if 'per person' in title else 'Minutes per participant'
                else:unit='pct';measure='Participation rate'
                # Counts printed below survey tables are controls, not activity estimates.
                if re.search(r'est(?:d|imated).*persons|sample.*persons|no\. of.*persons',first,re.I):
                    unit='hundreds' if '(00)' in first else 'count';measure='Source population/sample control'
                dims={**ctx,'source_row':first,'source_column':text,'measure':measure}
                # t1/t2/t45/t46 address geographic rows directly.
                if table['id'] in ['t1','t2','t45','t46']:dims['state']=first
                table['records'].append({'sheet':sheet.title,'cell':sheet.cell(ri,ci).coordinate,'raw':v,'value':value,'unit':unit,'dims':dims,'headers':path})
                roles[ri,ci]='data' if value is not None else 'unpublished'
                roles[ri,1]='row-label'
        populated=[];unassigned_numeric=[]
        for row in sheet:
            for cell in row:
                if cell.value is None:continue
                role=roles.get((cell.row,cell.column),'note-or-formatting')
                populated.append((cell.coordinate,role))
                if numeric(cell.value) is not None and role=='note-or-formatting':unassigned_numeric.append({'cell':cell.coordinate,'value':cell.value})
        if unassigned_numeric:raise ValueError(f'Unaccounted numeric source cells {table["id"]}: {unassigned_numeric[:5]}')
        table['sheets'].append({'name':sheet.title,'populatedCells':len(populated),'roles':{r:sum(x[1]==r for x in populated) for r in sorted(set(x[1] for x in populated))},'unassignedNumeric':unassigned_numeric})
    if not table['records']:raise ValueError('No interpreted cells '+table['id'])
    result['tables'].append(table)
    print(table['id'],len(table['records']),flush=True)
OUT.write_text(json.dumps(result,separators=(',',':'),ensure_ascii=False)+'\n')
